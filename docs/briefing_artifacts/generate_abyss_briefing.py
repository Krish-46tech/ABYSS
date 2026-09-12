from __future__ import annotations

import textwrap
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "briefing_artifacts"
DOCX_PATH = OUT_DIR / "ABYSS_Project_Briefing_Simple_A_to_Z.docx"


def font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Helvetica Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Helvetica.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
    ]
    for item in candidates:
        if Path(item).exists():
            return ImageFont.truetype(item, size)
    return ImageFont.load_default()


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        trial = f"{line} {word}".strip()
        if draw.textbbox((0, 0), trial, font=fnt)[2] <= width:
            line = trial
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def rounded_box(draw, xy, fill, outline, radius=18, width=3):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def draw_centered_text(draw, box, title, body="", title_size=30, body_size=20):
    x1, y1, x2, y2 = box
    title_font = font(title_size, True)
    body_font = font(body_size)
    max_width = x2 - x1 - 34
    lines = wrap(draw, title, title_font, max_width)
    body_lines = wrap(draw, body, body_font, max_width) if body else []
    total_h = len(lines) * (title_size + 6) + len(body_lines) * (body_size + 6)
    y = y1 + ((y2 - y1) - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=title_font)
        draw.text((x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2, y), line, font=title_font, fill="#0B1320")
        y += title_size + 6
    y += 2
    for line in body_lines:
        bbox = draw.textbbox((0, 0), line, font=body_font)
        draw.text((x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2, y), line, font=body_font, fill="#243447")
        y += body_size + 6


def arrow(draw, start, end, color="#334155", width=5):
    draw.line([start, end], fill=color, width=width)
    sx, sy = start
    ex, ey = end
    if ex >= sx:
        points = [(ex, ey), (ex - 18, ey - 10), (ex - 18, ey + 10)]
    else:
        points = [(ex, ey), (ex + 18, ey - 10), (ex + 18, ey + 10)]
    draw.polygon(points, fill=color)


def save_pipeline_diagram(path: Path):
    img = Image.new("RGB", (1800, 920), "#F8FAFC")
    draw = ImageDraw.Draw(img)
    draw.text((70, 44), "ABYSS A to Z Pipeline", font=font(42, True), fill="#0B1320")
    draw.text((70, 104), "Simple idea: sonar picture goes in, a prioritized detection report comes out.", font=font(25), fill="#334155")

    boxes = [
        ((80, 220, 340, 390), "#D9F99D", "1 Input", "Side scan sonar image"),
        ((430, 220, 690, 390), "#BAE6FD", "2 Clean", "Denoise, normalize, resize"),
        ((780, 220, 1040, 390), "#FDE68A", "3 Detect", "YOLOv8 finds boxes"),
        ((1130, 220, 1390, 390), "#DDD6FE", "4 Trust Score", "Shadow plus quality plus model score"),
        ((1480, 220, 1740, 390), "#FECACA", "5 Backend", "API returns JSON"),
        ((780, 560, 1040, 730), "#CFFAFE", "6 Frontend", "Upload, view boxes, inspect scores"),
        ((1130, 560, 1390, 730), "#BBF7D0", "7 Priority", "Rank what to check first"),
    ]
    for box, fill, title, body in boxes:
        rounded_box(draw, box, fill, "#CBD5E1")
        draw_centered_text(draw, box, title, body)
    for i in range(4):
        arrow(draw, (boxes[i][0][2] + 12, 305), (boxes[i + 1][0][0] - 12, 305))
    arrow(draw, (1610, 405), (1610, 480))
    arrow(draw, (1610, 480), (1045, 645))
    arrow(draw, (1055, 645), (1120, 645))
    draw.text((80, 820), "Think of it like a smart underwater teacher: it looks at a sonar image, points to suspicious objects, and says how sure it is.", font=font(25, True), fill="#0F172A")
    img.save(path, quality=95)


def save_backend_diagram(path: Path):
    img = Image.new("RGB", (1800, 960), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    draw.text((70, 44), "Backend Framework", font=font(42, True), fill="#0B1320")
    draw.text((70, 104), "FastAPI is the waiter between the frontend and the AI model.", font=font(25), fill="#334155")
    center = (760, 250, 1040, 430)
    rounded_box(draw, center, "#E0F2FE", "#0284C7")
    draw_centered_text(draw, center, "FastAPI Backend", "Receives requests and sends answers")
    endpoints = [
        ((90, 250, 370, 430), "#DCFCE7", "/health", "Is backend alive?"),
        ((90, 560, 370, 740), "#FEF3C7", "/detect", "Run model on uploaded image"),
        ((760, 560, 1040, 740), "#F3E8FF", "/geolocate", "Turn image pixel into map coordinate"),
        ((1430, 560, 1710, 740), "#FFE4E6", "/priority", "Rank detections by importance"),
        ((1430, 250, 1710, 430), "#E2E8F0", "Model Files", "Weights plus calibration JSON"),
    ]
    for box, fill, title, body in endpoints:
        rounded_box(draw, box, fill, "#CBD5E1")
        draw_centered_text(draw, box, title, body)
    arrow(draw, (380, 340), (750, 340))
    arrow(draw, (380, 650), (750, 385))
    arrow(draw, (900, 550), (900, 440))
    arrow(draw, (1420, 650), (1050, 385))
    arrow(draw, (1420, 340), (1050, 340))
    draw.text((90, 842), "Important honesty: geolocation is real only when real navigation metadata is supplied. Without it, the backend clearly labels the coordinate as simulated.", font=font(25, True), fill="#0F172A")
    img.save(path, quality=95)


def save_frontend_diagram(path: Path):
    img = Image.new("RGB", (1800, 900), "#F8FAFC")
    draw = ImageDraw.Draw(img)
    draw.text((70, 44), "Frontend Framework", font=font(42, True), fill="#0B1320")
    draw.text((70, 104), "React is the control room screen for the operator.", font=font(25), fill="#334155")
    boxes = [
        ((100, 250, 390, 445), "#DBEAFE", "Upload Image", "Operator selects a sonar tile"),
        ((500, 250, 790, 445), "#CFFAFE", "Detection Canvas", "Shows the sonar image and boxes"),
        ((900, 250, 1190, 445), "#EDE9FE", "Score Panel", "Shows confidence, quality, and shadow scores"),
        ((1300, 250, 1590, 445), "#DCFCE7", "3D View", "Shows contacts on a seabed scene"),
        ((700, 580, 1100, 755), "#FEF3C7", "Priority Table", "Sorts detections so humans check the most important first"),
    ]
    for box, fill, title, body in boxes:
        rounded_box(draw, box, fill, "#CBD5E1")
        draw_centered_text(draw, box, title, body)
    for i in range(3):
        arrow(draw, (boxes[i][0][2] + 12, 347), (boxes[i + 1][0][0] - 12, 347))
    arrow(draw, (990, 455), (900, 570))
    draw.text((100, 812), "It is already a working prototype, but the audit notes that the final Phase 5 dashboard is not formally signed off yet.", font=font(25, True), fill="#0F172A")
    img.save(path, quality=95)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold=False, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    header_tr_pr.append(tbl_header)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.08
            if row_idx == 0:
                set_cell_shading(cell, "1F4E79")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.bold = True
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F3F6FA")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True, "FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
    style_table(table)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_picture(doc, path: Path, caption: str, width=6.7):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cap.runs:
        r.italic = True
        r.font.size = Pt(9)
    doc.add_paragraph()


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.6)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(26)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for name, size in [("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)]:
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(0, 0, 0)


def build_doc():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    pipeline_png = OUT_DIR / "abyss_pipeline_simple.png"
    backend_png = OUT_DIR / "abyss_backend_framework.png"
    frontend_png = OUT_DIR / "abyss_frontend_framework.png"
    save_pipeline_diagram(pipeline_png)
    save_backend_diagram(backend_png)
    save_frontend_diagram(frontend_png)

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("ABYSS Project Briefing Simple A to Z")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(26)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("AI powered side scan sonar detection, backend, frontend, and pipeline explanation").bold = True
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.add_run("Prepared from the current repository on 12 September 2026").italic = True
    add_picture(doc, pipeline_png, "Figure 1. The whole ABYSS idea in one picture.", width=6.65)
    add_para(doc, "This briefing explains the project in very simple language so it can be presented confidently. Imagine ABYSS as a smart underwater helper. You give it a sonar picture. It cleans the picture, looks for objects like planes, ships, and shipwrecks, gives each finding a confidence score, and shows the results in a dashboard.")
    add_para(doc, "The current honest status is: Phase 0 to Phase 4 have been audited for the ML and backend work. A React frontend prototype exists and works with the backend, but the project notes say the final Phase 5 dashboard is not formally audited yet.")

    doc.add_page_break()

    add_heading(doc, "One Minute Presentation Story")
    add_para(doc, "ABYSS helps people inspect underwater sonar images faster. Sonar images are like blurry underwater shadows. A human can inspect them, but it takes time. ABYSS uses an AI detector and extra trust checks to point out suspicious objects and rank what should be checked first.")
    add_bullets(doc, [
        "Input: a side scan sonar image.",
        "Processing: clean and resize the image to 640 by 640 pixels.",
        "Detection: YOLOv8 draws boxes around possible objects.",
        "Confidence: the backend combines model confidence, image quality, and sonar shadow clues.",
        "Output: the dashboard shows boxes, scores, coordinates, and priority ranking.",
        "Human role: the system helps the operator decide what to inspect first; it does not replace expert review.",
    ])

    add_heading(doc, "A to Z What We Have Done")
    rows = [
        ["A", "Aim", "Build a sonar object detection system for marine debris and underwater anomalies."],
        ["B", "Base project", "Set up Python ML code, FastAPI backend, React frontend, tests, requirements, and project docs."],
        ["C", "Collected data structure", "Prepared raw, split, processed, experiment, and log folders for sonar datasets."],
        ["D", "Data ingestion", "Scripts copy datasets, read labels, count classes, and make manifests."],
        ["E", "Data splitting", "Images are split into train, validation, and test by inferred group so the same group does not appear in two splits."],
        ["F", "Preprocessing", "Images are denoised, normalized, and letterboxed to 640 pixels for YOLO."],
        ["G", "Detector", "YOLOv8 detects Plane, Ship, and Shipwreck classes."],
        ["H", "Better checkpoint", "The current deployed checkpoint is improved_v2_baseline_finetune_adamw.pt."],
        ["I", "Shadow features", "The system checks whether a detected object has a reasonable dark sonar shadow nearby."],
        ["J", "Image quality", "The system calculates if the image area is clear enough to trust."],
        ["K", "Calibration", "A calibration file turns raw scores into a more interpretable final probability."],
        ["L", "Evaluation", "We measured held out test metrics, FROC, ECE calibration, precision, recall, F1, and confusion."],
        ["M", "Experiments", "We tried bounded retrains for shipwreck loss, shipwreck augmentation, box weight, and mosaic schedule."],
        ["N", "Decision", "We kept the baseline 640 v4 model as deployed because experiments had tradeoffs."],
        ["O", "Backend", "FastAPI serves health, detect, geolocate, and priority endpoints."],
        ["P", "Frontend", "React dashboard uploads sonar images, shows boxes, displays score bars, geolocates detections, and ranks priorities."],
        ["Q", "Quality honesty", "We removed a filename based Plane override so the backend reports native model output honestly."],
    ]
    add_table(doc, ["Letter", "Topic", "Simple explanation"], rows, [0.45, 1.35, 4.75])

    add_heading(doc, "Pipeline Explained Like You Are 10")
    add_picture(doc, pipeline_png, "Figure 2. The ABYSS pipeline from input image to operator decision.", width=6.7)
    add_para(doc, "Think of the pipeline like washing, reading, checking, and reporting a worksheet.")
    add_bullets(doc, [
        "First, ABYSS gets a sonar image. A sonar image is made from sound, not normal light.",
        "Second, it cleans the image because sonar pictures can be noisy.",
        "Third, YOLOv8 looks at the image and draws boxes around likely objects.",
        "Fourth, ABYSS asks extra questions: Is the picture clear? Does the object have a shadow that makes sense for sonar?",
        "Fifth, the backend sends a neat JSON answer to the frontend.",
        "Sixth, the frontend shows the result in a way a human operator can understand.",
    ])

    add_heading(doc, "Backend Complete Idea")
    add_picture(doc, backend_png, "Figure 3. Backend endpoints and how they fit together.", width=6.7)
    add_para(doc, "The backend is the brain room. It is written with FastAPI. The frontend asks it questions, and the backend answers with JSON data.")
    add_table(doc, ["Endpoint", "What it does", "Simple example"], [
        ["/health", "Checks whether the backend is alive.", "The frontend shows Backend reachable or Backend unavailable."],
        ["/detect", "Receives an uploaded image and runs YOLOv8 plus confidence scoring.", "Returns boxes, class names, confidence scores, and latency."],
        ["/geolocate", "Converts the middle of a detection box into latitude and longitude.", "Real if navigation metadata is supplied; simulated fallback otherwise."],
        ["/priority", "Sorts detections by a transparent formula.", "Uses confidence, hazard class weight, object size, and optional distance to sensitive zone."],
    ], [1.15, 3.05, 2.25])
    add_para(doc, "The backend loads two important files: the trained model weights and the calibration JSON. It checks that the calibration file matches the model, image size, preprocessing mode, and model hash. This matters because using the wrong calibration with the wrong model would make the confidence numbers untrustworthy.")
    add_para(doc, "Detection flow inside the backend:", bold_lead="")
    add_bullets(doc, [
        "Read the uploaded image bytes and decode them with OpenCV.",
        "Clean and resize the image using the same baseline preprocessing used during evaluation.",
        "Run YOLOv8 on the 640 pixel image.",
        "For each box, calculate shadow features and image quality.",
        "Build a three number feature vector: detector confidence, shadow consistency, image quality.",
        "Use logistic fusion and temperature calibration to produce composite confidence.",
        "Return a list of detections with class name, box coordinates, raw confidence, shadow score, quality score, fused probability, calibrated probability, and composite confidence.",
    ])
    add_table(doc, ["Backend score", "Meaning in simple language"], [
        ["detector_confidence", "How sure YOLO is that the object exists."],
        ["image_quality_score", "How clear and useful the image area looks."],
        ["shadow_consistency_score", "Whether the dark sonar shadow looks believable."],
        ["logistic_fused_probability", "A learned combination of the three clues."],
        ["temperature_calibrated_probability", "The final adjusted probability used as composite confidence."],
        ["priority_score", "Operational ranking score, not the same as detection probability."],
    ], [2.15, 4.3])

    add_heading(doc, "Frontend Complete Idea")
    add_picture(doc, frontend_png, "Figure 4. Frontend operator workflow.", width=6.7)
    add_para(doc, "The frontend is the control room screen. It is built with React, TypeScript, Vite, Tailwind CSS, lucide icons, and React Three Fiber for the 3D scene.")
    add_table(doc, ["Frontend part", "What the user sees", "What it talks to"], [
        ["Upload form", "Choose a sonar image and press Detect.", "/detect"],
        ["Sonar canvas", "Image with detection boxes and labels drawn on top.", "Uses /detect response"],
        ["Confidence breakdown", "Bar charts for detector, image quality, shadow, and composite confidence.", "Uses /detect response"],
        ["3D contact view", "Simple seabed scene with markers for detections.", "/geolocate"],
        ["Priority table", "Ranked list of what to inspect first.", "/priority"],
        ["Backend status", "Reachable or unavailable indicator.", "/health"],
    ], [1.75, 2.75, 1.65])
    add_para(doc, "In presentation language: the frontend does not do the AI thinking itself. It sends the image to the backend, receives the answer, and makes the answer easy for a human to inspect.")

    add_heading(doc, "Machine Learning Framework")
    add_para(doc, "The ML framework is centered on YOLOv8. YOLO means You Only Look Once. In simple words, it looks at the image once and predicts boxes and classes quickly.")
    add_table(doc, ["Layer", "Tool or file", "Simple purpose"], [
        ["Data", "abyss/data and data.yaml", "Stores train, validation, and test images and labels."],
        ["Preprocessing", "preprocess.py", "Makes images cleaner and the same size."],
        ["Training", "train_yolo.py and train_bounded_candidate.py", "Creates or experiments with model checkpoints."],
        ["Inference", "infer.py and backend services.py", "Runs the model on new images."],
        ["Calibration", "fit_calibration.py and calibration_improved_v4.json", "Makes confidence scores easier to trust."],
        ["Evaluation", "evaluate.py, evaluate_froc.py, evaluate_calibration.py", "Measures accuracy, false alarms, and calibration."],
    ], [1.45, 2.35, 2.65])
    add_para(doc, "The classes are Plane, Ship, and Shipwreck. The model is best at Ship, decent at Plane by AP, and currently weak on Shipwreck. That weakness should be mentioned honestly.")

    add_heading(doc, "What The Results Say")
    add_table(doc, ["Class", "Ground truth objects", "AP at IoU 0.5", "AP 0.5 to 0.95", "Simple meaning"], [
        ["Plane", "17", "0.779424", "0.378747", "Often finds planes, but many are confused with ships at the deployed threshold."],
        ["Ship", "74", "0.957086", "0.674665", "Strongest class."],
        ["Shipwreck", "20", "0.144992", "0.058058", "Weak class and needs more work."],
        ["Overall", "111", "0.627168", "0.370490", "Useful prototype, not perfect production truth."],
    ], [1.05, 1.2, 1.05, 1.15, 2.0])
    add_para(doc, "At the deployed raw confidence threshold of 0.25, the held out test results were 77 true positives, 44 false positives, and 34 false negatives. Precision was 0.636364, recall was 0.693694, and F1 was 0.663793.")
    add_picture(doc, ROOT / "docs" / "figures" / "experiment_overall_map.png", "Figure 5. Overall held out mAP comparison for baseline and bounded experiments.", width=6.35)
    add_picture(doc, ROOT / "docs" / "figures" / "experiment_class_ap50.png", "Figure 6. Per class held out AP at IoU 0.5.", width=6.35)
    add_picture(doc, ROOT / "docs" / "figures" / "baseline_ece_comparison.png", "Figure 7. Calibration comparison. Lower ECE means confidence is better aligned with correctness.", width=6.35)
    add_picture(doc, ROOT / "docs" / "figures" / "experiment_decision_flow.png", "Figure 8. Why the deployed model was kept instead of switching to an experimental candidate.", width=6.35)

    add_heading(doc, "Experiment Story")
    add_para(doc, "After the baseline, we tried careful bounded experiments. Bounded means we changed one idea at a time so we could understand what helped or hurt.")
    add_table(doc, ["Candidate", "Goal", "Outcome"], [
        ["Shipwreck loss weight", "Make the model care more about shipwrecks.", "Shipwreck improved slightly, but overall validation and some class tradeoffs were not good enough."],
        ["Shipwreck only augmentation", "Create more shipwreck examples.", "Shipwreck improved, but Plane and overall results dropped."],
        ["Box weight 9.0", "Encourage better box placement.", "Some AP improved, but main target mAP 0.5 to 0.95 dropped."],
        ["Close mosaic 15", "Change mosaic augmentation schedule.", "Small aggregate gain, but Plane and Ship AP 0.5 fell, so kept as experimental only."],
    ], [1.8, 2.25, 2.4])
    add_para(doc, "Final decision: keep baseline 640 v4 deployed. Any future model change must get a new calibration file and a full rerun of corrected metrics.")

    add_heading(doc, "Important Limitations To Say Clearly")
    add_bullets(doc, [
        "Shipwreck detection is still weak and needs more data, better training, or a better strategy.",
        "The model can confuse Plane and Ship. Earlier code had a filename based Plane override, but it has been removed so results are honest.",
        "Geolocation is real only when survey navigation metadata is provided. Without it, the backend uses a simulated fallback origin and labels it as simulated.",
        "We cannot calculate FP per kilometer, FP per square kilometer, or positional RMSE because real survey distance, area, sensor specs, and GPS ground truth are missing.",
        "Priority score is a transparent policy formula, not a learned danger probability.",
        "The React dashboard works as a prototype, but the audit says final Phase 5 dashboard work has not been formally signed off.",
    ])

    add_heading(doc, "How To Explain Frameworks In The Presentation")
    add_table(doc, ["Framework", "Child friendly explanation", "Why ABYSS uses it"], [
        ["FastAPI", "A fast waiter that takes requests and brings answers.", "Serves detection, geolocation, priority, and health APIs."],
        ["React", "A way to build a screen from small reusable blocks.", "Builds the operator dashboard."],
        ["Vite", "A quick starter engine for the web app.", "Runs the frontend during development."],
        ["Tailwind CSS", "Small ready made style rules.", "Makes the dashboard clean without writing huge CSS files."],
        ["React Three Fiber", "React for 3D scenes.", "Shows the seabed contact view."],
        ["OpenCV", "A computer vision toolbox.", "Reads, cleans, resizes, and measures sonar images."],
        ["YOLOv8", "An AI model that draws boxes around objects.", "Detects Plane, Ship, and Shipwreck."],
        ["Pydantic", "A strict form checker.", "Validates API request and response shapes."],
    ], [1.35, 2.75, 2.35])

    add_heading(doc, "Suggested Spoken Script")
    add_para(doc, "Here is a simple way to present the project:")
    script = [
        "Our project is called ABYSS. It helps inspect side scan sonar images.",
        "Side scan sonar images are not normal photos. They are made from sound, so they can be noisy and hard to read.",
        "First we clean the image and resize it to 640 by 640 pixels.",
        "Then YOLOv8 looks for three object types: Plane, Ship, and Shipwreck.",
        "After that, we do not blindly trust the model. We also check image quality and sonar shadow clues.",
        "The backend combines those clues into a final confidence score and sends the result as JSON.",
        "The frontend lets a user upload an image, see boxes on the sonar image, inspect confidence bars, see a 3D contact view, and read a priority list.",
        "Our current model is strongest on Ship, okay on Plane by AP, and still weak on Shipwreck. We are honest about that because this is an engineering project, not magic.",
        "The next best work is improving shipwreck data and retraining with proper calibration.",
    ]
    for line in script:
        add_para(doc, line)

    add_heading(doc, "Quick Q and A")
    add_table(doc, ["Question", "Good answer"], [
        ["What problem are we solving?", "We help humans quickly find suspicious underwater objects in sonar images."],
        ["Is it fully automatic?", "No. It is a decision support tool. Humans still verify important detections."],
        ["What does backend do?", "It runs detection, scoring, geolocation, and priority ranking."],
        ["What does frontend do?", "It makes backend answers visible and understandable."],
        ["What is the pipeline?", "Input sonar image, clean image, detect objects, calculate trust scores, serve API result, show dashboard."],
        ["What is the biggest weakness?", "Shipwreck detection and lack of real navigation metadata."],
        ["What is the biggest strength?", "The project is end to end: data pipeline, detector, calibration, backend API, and frontend prototype."],
    ], [2.15, 4.25])

    add_heading(doc, "Project File Map")
    add_table(doc, ["Folder or file", "What it contains"], [
        ["README.md", "High level architecture, roadmap, quickstart, and API overview."],
        ["abyss/ml/preprocessing", "Data ingestion, splitting, denoising, normalizing, resizing, label conversion."],
        ["abyss/ml/detection", "YOLO training, inference, evaluation, bounded experiments, model weights."],
        ["abyss/ml/shadow_confidence", "Shadow features, image quality, calibration, ECE, FROC, threshold checks."],
        ["abyss/backend/app", "FastAPI routes, schemas, and backend service logic."],
        ["frontend/src/App.tsx", "Main React dashboard prototype."],
        ["docs/METRICS.md", "Canonical current deployed metrics."],
        ["docs/PROJECT_METRICS.md", "Combined work summary, charts, and experimental comparison."],
        ["logs", "Saved experiment reports and audit outputs."],
    ], [2.2, 4.25])

    add_heading(doc, "Final Summary")
    add_para(doc, "ABYSS is an end to end sonar detection project. It is not just a model file. It includes data preparation, model training, confidence calibration, backend APIs, evaluation reports, experiment audits, and a frontend prototype. The best way to present it is as a human assist tool: it helps an operator find and prioritize suspicious underwater objects faster.")
    add_para(doc, "The correct final message is confident but honest: we built a working pipeline and prototype, measured it carefully, kept the safest deployed model, and documented what still needs work.")

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_doc()
